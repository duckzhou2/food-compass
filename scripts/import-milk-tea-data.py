"""Convert the reviewed drink-brand workbooks into stable frontend JSON.

The source workbooks are read-only inputs. Missing fields stay null, lower-priority
records never overwrite a more specific combination, and source-side review
decisions are copied only to the internal import report.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import warnings
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


BRANDS = {
    "一点点.xlsx": ("yidiandian", "一点点", "产品名称", "标准分类"),
    "喜茶.xlsx": ("heytea", "喜茶", "产品名称", "分类"),
    "蜜雪冰城.xlsx": ("mixue", "蜜雪冰城", "产品名称", "分类"),
    "霸王茶姬.xlsx": ("chagee", "霸王茶姬", "产品名称", "分类"),
    "CoCo都可.xlsx": ("coco", "CoCo都可", "产品名称", "原始分类"),
    "古茗.xlsx": ("guming", "古茗", "产品名称", "原始分类"),
    "瑞幸咖啡.xlsx": ("luckin", "瑞幸咖啡", "产品名称", "原始分类"),
    "茶百道.xlsx": ("chabaidao", "茶百道", "产品名称", "原始分类"),
}

SELECTION_FIELDS = ("size", "ice", "temperature", "sugar", "version", "base")
STATUS_ORDER = {"missing": 0, "unspecified_reference": 1, "specified_reference": 2, "detailed": 3}
UNKNOWN_VALUES = {"", "未说明", "规格未说明", "不详", "未知"}
DISPLAY_CATEGORIES = (
    "奶茶 / 奶绿",
    "鲜奶茶 / 轻乳茶",
    "纯茶",
    "果茶 / 果饮",
    "茶特调",
    "冰淇淋 / 甜品",
    "咖啡",
)


def stable_id(prefix: str, *parts: object) -> str:
    payload = "\u241f".join("" if part is None else str(part).strip() for part in parts)
    digest = hashlib.sha1(payload.encode("utf-8")).hexdigest()[:12]
    return f"{prefix}-{digest}"


def clean_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return None if text in UNKNOWN_VALUES else text


def numeric(value: object) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    return None


def rows_from_sheet(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    worksheet = workbook[sheet_name]
    row_iterator = worksheet.iter_rows(values_only=True)
    headers = next(row_iterator)
    result: list[dict[str, Any]] = []
    for row in row_iterator:
        if not any(value is not None for value in row):
            continue
        result.append({str(headers[index]): value for index, value in enumerate(row) if headers[index] is not None})
    return result


def split_ice_temperature(value: object) -> tuple[str | None, str | None]:
    text = clean_text(value)
    if text in {"热", "温", "常温"}:
        return None, text
    return text, None


def parse_default_spec(value: object) -> dict[str, str | None]:
    selection = {field: None for field in SELECTION_FIELDS}
    text = clean_text(value)
    if text is None:
        return selection
    tokens = [token.strip() for token in text.replace("；", "·").split("·") if token.strip()]
    for token in tokens:
        if any(marker in token for marker in ("简称", "分开记录", "原图")):
            continue
        if token in {"热", "温", "常温"}:
            selection["temperature"] = token
        elif "糖" in token or token.endswith("甜"):
            selection["sugar"] = token
        elif "冰" in token or "沙" in token:
            selection["ice"] = token.replace("正常冰", "正常冰")
        elif token.endswith("杯") or token in {"瓶装", "杯", "份"}:
            selection["size"] = token
        elif "气泡" in token:
            selection["version"] = token
        elif "版" in token or "基底" in token:
            selection["base"] = token
    return selection


def data_status(selection: dict[str, Any], calories: int | float | None) -> str:
    if calories is None:
        return "missing"
    specificity = sum(selection[field] is not None for field in SELECTION_FIELDS)
    if specificity >= 2:
        return "detailed"
    if specificity == 1:
        return "specified_reference"
    return "unspecified_reference"


def display_category(raw_category: str, product_name: str) -> str:
    text = raw_category.strip()
    if "冰淇淋" in text or "甜品" in text or product_name == "雪王冰杯":
        return "冰淇淋 / 甜品"
    if "咖啡" in text or "茶咖" in text:
        return "咖啡"
    if "纯茶" in text or "清饮" in text:
        return "纯茶"
    if "特调" in text:
        return "茶特调"
    if "轻乳" in text or "鲜奶" in text or ("牛乳" in text and "果" not in text):
        return "鲜奶茶 / 轻乳茶"
    if any(marker in text for marker in ("奶茶", "奶绿", "奶乌", "奶青", "乳饮", "撞奶", "酸奶", "巧克力", "可可", "麦芽")):
        return "奶茶 / 奶绿"
    if any(marker in text for marker in ("果", "柠檬", "乳酸菌", "椰饮")):
        return "果茶 / 果饮"
    if any(marker in text for marker in ("抹茶", "奶盖", "招牌茶饮", "茗茶")):
        return "茶特调"
    return "茶特调"


def make_variant(
    brand_id: str,
    product_name: str,
    selection: dict[str, str | None],
    calories: int | float | None,
    notes: Iterable[object] = (),
    review_flags: Iterable[str] = (),
) -> dict[str, Any]:
    variant_id = stable_id(
        f"{brand_id}-variant",
        product_name,
        *(selection[field] for field in SELECTION_FIELDS),
    )
    flags = sorted(set(review_flags))
    return {
        "variantId": variant_id,
        **selection,
        "calories": calories,
        "dataStatus": "needs_review" if flags else data_status(selection, calories),
        "notes": [text for value in notes if (text := clean_text(value)) is not None],
        "reviewFlags": flags,
    }


def selection_key(variant: dict[str, Any]) -> tuple[Any, ...]:
    return tuple(variant[field] for field in SELECTION_FIELDS)


def selection_specificity(variant: dict[str, Any]) -> int:
    return sum(variant[field] is not None for field in SELECTION_FIELDS)


def choose_default(variants: list[dict[str, Any]], explicit_id: str | None = None) -> dict[str, Any]:
    if explicit_id is not None:
        explicit = next((variant for variant in variants if variant["variantId"] == explicit_id), None)
        if explicit is not None and explicit["dataStatus"] != "needs_review" and explicit["calories"] is not None:
            return explicit

    normal_variants = [variant for variant in variants if variant["dataStatus"] != "needs_review"]
    candidates = normal_variants or variants
    size_rank = {"中杯": 0, "标准杯": 0, "主题杯": 0, "小杯": 1, "大杯": 2}
    return sorted(
        candidates,
        key=lambda variant: (
            variant["calories"] is None,
            size_rank.get(variant["size"], 1),
            -selection_specificity(variant),
            variant["base"] not in (None, "A2牛乳版"),
            variant["version"] is not None,
            variant["variantId"],
        ),
    )[0]


def default_selection(variant: dict[str, Any]) -> dict[str, Any]:
    return {"variantId": variant["variantId"], **{field: variant[field] for field in SELECTION_FIELDS}}


def product_record(
    brand_id: str,
    brand_name: str,
    catalog_row: dict[str, Any],
    name_field: str,
    category_field: str,
    variants: list[dict[str, Any]],
    topping_ids: list[str],
    explicit_default_id: str | None = None,
    review_flags: Iterable[str] = (),
) -> dict[str, Any]:
    product_name = str(catalog_row[name_field]).strip()
    if not variants:
        variants = [make_variant(brand_id, product_name, {field: None for field in SELECTION_FIELDS}, None)]
    default = choose_default(variants, explicit_default_id)
    product_flags = sorted(set(review_flags))
    if default["dataStatus"] == "needs_review" and not any(
        variant["dataStatus"] != "needs_review" and variant["calories"] is not None
        for variant in variants
    ):
        product_flags = sorted(set([*product_flags, "review-only-product"]))
    raw_category = clean_text(catalog_row.get(category_field)) or "未分类"
    return {
        "brandId": brand_id,
        "brandName": brand_name,
        "productId": stable_id(brand_id, brand_name, product_name),
        "productName": product_name,
        "category": raw_category,
        "displayCategory": display_category(raw_category, product_name),
        "defaultSelection": default_selection(default),
        "variants": variants,
        "toppings": topping_ids,
        "notes": [text for value in (catalog_row.get("备注"),) if (text := clean_text(value)) is not None],
        "dataStatus": default["dataStatus"],
        "reviewFlags": product_flags,
    }


def make_topping_variant(
    brand_id: str,
    topping_name: str,
    size: str | None,
    unit: str | None,
    minimum: int | float | None,
    maximum: int | float | None,
    notes: Iterable[object] = (),
) -> dict[str, Any]:
    calories = None if minimum is None or maximum is None else {"min": minimum, "max": maximum}
    selection = {field: None for field in SELECTION_FIELDS}
    selection["size"] = size
    return {
        "toppingVariantId": stable_id(f"{brand_id}-topping-variant", topping_name, size, unit),
        "size": size,
        "unit": unit,
        "calories": calories,
        "dataStatus": data_status(selection, minimum),
        "notes": [text for value in notes if (text := clean_text(value)) is not None],
    }


def topping_spec(value: object) -> tuple[str | None, str | None]:
    unit = clean_text(value)
    if unit is None:
        return None, None
    size = None
    if "/" in unit:
        possible_size = unit.split("/", 1)[0].strip()
        if possible_size.endswith("杯"):
            size = possible_size
    return size, unit


def group_toppings(brand_id: str, variants_by_name: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    result = []
    for name in sorted(variants_by_name, key=lambda value: value.encode("utf-8")):
        variants = variants_by_name[name]
        specific = [variant for variant in variants if variant["size"] is not None or variant["unit"] is not None]
        if specific:
            variants = specific
        unique = {variant["toppingVariantId"]: variant for variant in variants}
        result.append({"toppingId": stable_id(f"{brand_id}-topping", name), "name": name, "variants": list(unique.values())})
    return result


def build_yidiandian(workbook: Any, report: dict[str, Any]) -> dict[str, Any]:
    brand_id, brand_name, name_field, category_field = BRANDS["一点点.xlsx"]
    catalog = rows_from_sheet(workbook, "产品目录")
    heat_rows = rows_from_sheet(workbook, "热量记录")
    candidates: dict[str, list[tuple[dict[str, Any], int, str]]] = defaultdict(list)

    for row in heat_rows:
        product_name = str(row[name_field]).strip()
        ice, temperature = split_ice_temperature(None)
        selection = {
            "size": clean_text(row.get("杯型")),
            "ice": ice,
            "temperature": temperature,
            "sugar": None,
            "version": clean_text(row.get("定制条件")),
            "base": clean_text(row.get("奶底/版本")),
        }
        source_id = clean_text(row.get("来源编号")) or ""
        source_rank = {"SRC-A": 1, "SRC-B": 2, "SRC-C": 3}.get(source_id, 0)
        variant = make_variant(
            brand_id,
            product_name,
            selection,
            numeric(row.get("参考热量(kcal)")),
            (row.get("口径说明"), row.get("备注")),
            ("conflicting-calorie-sources",) if product_name == "六窨香柠绿" else (),
        )
        candidates[product_name].append((variant, source_rank, str(row.get("记录ID") or "")))

    variants_by_product: dict[str, list[dict[str, Any]]] = {}
    for product_name, entries in candidates.items():
        maximum_specificity = max(selection_specificity(variant) for variant, _, _ in entries)
        for variant, _, record_id in entries:
            if selection_specificity(variant) < maximum_specificity:
                report["skippedRecords"].append(
                    {"brand": brand_name, "product": product_name, "recordId": record_id, "calories": variant["calories"], "reason": "存在规格信息更完整的记录组"}
                )
        entries = [entry for entry in entries if selection_specificity(entry[0]) == maximum_specificity]

        selected: dict[tuple[Any, ...], tuple[dict[str, Any], int, str]] = {}
        for entry in entries:
            variant, source_rank, record_id = entry
            key = selection_key(variant)
            current = selected.get(key)
            if current is None or (source_rank, record_id) > (current[1], current[2]):
                if current is not None:
                    report["skippedRecords"].append(
                        {"brand": brand_name, "product": product_name, "recordId": current[2], "calories": current[0]["calories"], "reason": "同规格冲突，保留较新的同口径记录"}
                    )
                selected[key] = entry
            else:
                report["skippedRecords"].append(
                    {"brand": brand_name, "product": product_name, "recordId": record_id, "calories": variant["calories"], "reason": "同规格冲突，保留较新的同口径记录"}
                )
        variants_by_product[product_name] = [entry[0] for entry in selected.values()]

    topping_candidates: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows_from_sheet(workbook, "小料热量记录"):
        name = str(row["小料名称"]).strip()
        size, unit = topping_spec(row.get("规格"))
        calories = numeric(row.get("参考热量(kcal)"))
        topping_candidates[name].append(
            make_topping_variant(brand_id, name, size, unit, calories, calories, (row.get("口径说明"), row.get("备注")))
        )
    toppings = group_toppings(brand_id, topping_candidates)
    topping_ids = [topping["toppingId"] for topping in toppings]
    products = [
        product_record(
            brand_id,
            brand_name,
            row,
            name_field,
            category_field,
            variants_by_product.get(str(row[name_field]).strip(), []),
            topping_ids,
        )
        for row in catalog
    ]
    return {"brandId": brand_id, "brandName": brand_name, "products": products, "toppings": toppings}


def heytea_matrix_selection(row: dict[str, Any]) -> dict[str, str | None]:
    ice, temperature = split_ice_temperature(row.get("温度/冰量"))
    specification = clean_text(row.get("规格说明"))
    size = None
    if specification:
        if "超大杯" in specification:
            size = "超大杯"
        elif "瓶装" in specification:
            size = "瓶装"
        elif "标准杯" in specification:
            size = "标准杯"
    return {
        "size": size,
        "ice": ice,
        "temperature": temperature,
        "sugar": clean_text(row.get("甜度")),
        "version": None,
        "base": None,
    }


def build_heytea(workbook: Any, report: dict[str, Any]) -> dict[str, Any]:
    brand_id, brand_name, name_field, category_field = BRANDS["喜茶.xlsx"]
    catalog = rows_from_sheet(workbook, "产品目录")
    heat_rows = rows_from_sheet(workbook, "热量矩阵")
    variants_by_product: dict[str, list[dict[str, Any]]] = defaultdict(list)
    defaults: dict[str, str] = {}
    for row in heat_rows:
        product_name = str(row[name_field]).strip()
        variant = make_variant(
            brand_id,
            product_name,
            heytea_matrix_selection(row),
            numeric(row.get("参考热量(kcal)")),
            (row.get("备注"),),
            ("source-note-review",) if "异常" in str(row.get("备注") or "") else (),
        )
        variants_by_product[product_name].append(variant)
        if clean_text(row.get("是否默认规格")) == "是":
            defaults[product_name] = variant["variantId"]

    for row in catalog:
        product_name = str(row[name_field]).strip()
        if product_name not in variants_by_product:
            selection = parse_default_spec(row.get("默认规格"))
            variants_by_product[product_name].append(
                make_variant(
                    brand_id,
                    product_name,
                    selection,
                    numeric(row.get("默认参考热量(kcal)")),
                    (row.get("备注"), "仅提供单一默认规格参考值"),
                )
            )
            defaults[product_name] = variants_by_product[product_name][0]["variantId"]

    topping_candidates: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows_from_sheet(workbook, "小料与云顶"):
        name = str(row["名称"]).strip()
        size, unit = topping_spec(row.get("规格"))
        topping_candidates[name].append(
            make_topping_variant(
                brand_id,
                name,
                size,
                unit,
                numeric(row.get("最低参考热量")),
                numeric(row.get("最高参考热量")),
                (row.get("备注"),),
            )
        )
    toppings = group_toppings(brand_id, topping_candidates)
    topping_ids = [topping["toppingId"] for topping in toppings]
    products = []
    for row in catalog:
        product_name = str(row[name_field]).strip()
        status = clean_text(row.get("数据状态")) or ""
        flags = ["catalog-mismatch"] if "差异" in status else []
        products.append(
            product_record(
                brand_id,
                brand_name,
                row,
                name_field,
                category_field,
                variants_by_product[product_name],
                topping_ids,
                defaults.get(product_name),
                flags,
            )
        )
    return {"brandId": brand_id, "brandName": brand_name, "products": products, "toppings": toppings}


def build_mixue(workbook: Any, report: dict[str, Any]) -> dict[str, Any]:
    brand_id, brand_name, name_field, category_field = BRANDS["蜜雪冰城.xlsx"]
    catalog = rows_from_sheet(workbook, "产品目录")
    heat_rows = rows_from_sheet(workbook, "热量矩阵")
    variants_by_product: dict[str, list[dict[str, Any]]] = defaultdict(list)
    defaults: dict[str, str] = {}
    heat_products: set[str] = set()
    for row in heat_rows:
        product_name = str(row[name_field]).strip()
        heat_products.add(product_name)
        ice, temperature = split_ice_temperature(row.get("温度/冰量"))
        selection = {
            "size": clean_text(row.get("杯型")),
            "ice": ice,
            "temperature": temperature,
            "sugar": clean_text(row.get("甜度")),
            "version": None,
            "base": None,
        }
        variant = make_variant(brand_id, product_name, selection, numeric(row.get("参考热量(kcal)")), (row.get("备注"),))
        variants_by_product[product_name].append(variant)
        if clean_text(row.get("是否默认规格")) == "是":
            defaults[product_name] = variant["variantId"]

    for row in catalog:
        product_name = str(row[name_field]).strip()
        if product_name in heat_products:
            continue
        ice, temperature = split_ice_temperature(row.get("温度/冰量"))
        selection = {
            "size": clean_text(row.get("杯型")),
            "ice": ice,
            "temperature": temperature,
            "sugar": None,
            "version": None,
            "base": None,
        }
        variant = make_variant(
            brand_id,
            product_name,
            selection,
            numeric(row.get("默认参考热量(kcal)")),
            (row.get("备注"), "原表仅提供该杯型/冰量的默认参考值，未展开甜度组合"),
        )
        variants_by_product[product_name].append(variant)
        if clean_text(row.get("默认加入转盘")) == "是":
            defaults[product_name] = variant["variantId"]

    topping_candidates: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows_from_sheet(workbook, "小料目录"):
        name = str(row["小料名称"]).strip()
        size, unit = topping_spec(row.get("规格"))
        calories = numeric(row.get("参考热量(kcal)"))
        topping_candidates[name].append(
            make_topping_variant(brand_id, name, size, unit, calories, calories, (row.get("备注"),))
        )
    toppings = group_toppings(brand_id, topping_candidates)
    topping_ids = [topping["toppingId"] for topping in toppings]

    catalog_by_name: dict[str, dict[str, Any]] = {}
    catalog_order: list[str] = []
    for row in catalog:
        product_name = str(row[name_field]).strip()
        if product_name not in catalog_by_name:
            catalog_order.append(product_name)
            catalog_by_name[product_name] = row
    products = [
        product_record(
            brand_id,
            brand_name,
            catalog_by_name[name],
            name_field,
            category_field,
            variants_by_product[name],
            topping_ids,
            defaults.get(name),
        )
        for name in catalog_order
    ]
    return {"brandId": brand_id, "brandName": brand_name, "products": products, "toppings": toppings}


def build_chagee(workbook: Any, report: dict[str, Any]) -> dict[str, Any]:
    brand_id, brand_name, name_field, category_field = BRANDS["霸王茶姬.xlsx"]
    catalog = rows_from_sheet(workbook, "产品目录")
    heat_rows = rows_from_sheet(workbook, "热量矩阵")
    variants_by_product: dict[str, list[dict[str, Any]]] = defaultdict(list)
    defaults: dict[str, str] = {}
    for row in heat_rows:
        product_name = str(row[name_field]).strip()
        ice, temperature = split_ice_temperature(row.get("冰量/温度"))
        version_or_base = clean_text(row.get("版本/基底"))
        selection = {
            "size": clean_text(row.get("杯型")),
            "ice": ice,
            "temperature": temperature,
            "sugar": clean_text(row.get("甜度")),
            "version": version_or_base if version_or_base and "气泡" in version_or_base else None,
            "base": version_or_base if version_or_base and "气泡" not in version_or_base else None,
        }
        grade = clean_text(row.get("数据等级")) or ""
        review_flags = ["source-marked-review"] if "C" in grade else []
        variant = make_variant(
            brand_id,
            product_name,
            selection,
            numeric(row.get("参考热量(kcal)")),
            (row.get("备注"),),
            review_flags,
        )
        variants_by_product[product_name].append(variant)
        if clean_text(row.get("是否默认规格")) == "是":
            defaults[product_name] = variant["variantId"]

    products = []
    for row in catalog:
        product_name = str(row[name_field]).strip()
        grade = clean_text(row.get("数据等级")) or ""
        flags = ["source-marked-review"] if "C" in grade else []
        products.append(
            product_record(
                brand_id,
                brand_name,
                row,
                name_field,
                category_field,
                variants_by_product[product_name],
                [],
                defaults.get(product_name),
                flags,
            )
        )
    return {"brandId": brand_id, "brandName": brand_name, "products": products, "toppings": []}


def build_normalized(workbook: Any, report: dict[str, Any], filename: str, heat_sheet: str) -> dict[str, Any]:
    """Import workbooks that already expose the shared product/variant/topping columns."""
    brand_id, brand_name, name_field, category_field = BRANDS[filename]
    catalog = rows_from_sheet(workbook, "产品目录")
    catalog_by_name = {str(row[name_field]).strip(): row for row in catalog}
    variants_by_product: dict[str, list[dict[str, Any]]] = defaultdict(list)
    explicit_defaults: dict[str, str] = {}
    selected_by_product: dict[str, dict[tuple[Any, ...], dict[str, Any]]] = defaultdict(dict)

    for row in rows_from_sheet(workbook, heat_sheet):
        product_name = str(row[name_field]).strip()
        ice, temperature = split_ice_temperature(row.get("饮用方式"))
        selection = {
            "size": clean_text(row.get("杯型")),
            "ice": ice,
            "temperature": temperature,
            "sugar": clean_text(row.get("甜度")),
            "version": clean_text(row.get("版本")),
            "base": clean_text(row.get("奶底或茶底")),
        }
        catalog_status = clean_text(catalog_by_name.get(product_name, {}).get("数据状态"))
        row_status = clean_text(row.get("数据状态"))
        flags = ("source-needs-review",) if "needs-review" in {catalog_status, row_status} else ()
        variant = make_variant(
            brand_id,
            product_name,
            selection,
            numeric(row.get("参考热量(kcal)")),
            (row.get("其他可选规格"), row.get("内部备注")),
            flags,
        )
        key = selection_key(variant)
        current = selected_by_product[product_name].get(key)
        is_default = str(row.get("默认规格") or "").startswith("是")
        if current is None or (is_default and current["calories"] != variant["calories"]):
            if current is not None:
                report["skippedRecords"].append({
                    "brand": brand_name,
                    "product": product_name,
                    "recordId": current["variantId"],
                    "calories": current["calories"],
                    "reason": "同规格冲突，保留原表标记的默认记录",
                })
            selected_by_product[product_name][key] = variant
        elif current["calories"] != variant["calories"]:
            report["skippedRecords"].append({
                "brand": brand_name,
                "product": product_name,
                "recordId": variant["variantId"],
                "calories": variant["calories"],
                "reason": "同规格冲突，保留更早的明确记录",
            })
        if is_default:
            explicit_defaults[product_name] = variant["variantId"]

    for product_name, selected in selected_by_product.items():
        variants_by_product[product_name] = list(selected.values())

    topping_candidates: dict[str, list[dict[str, Any]]] = defaultdict(list)
    if "小料目录" in workbook.sheetnames:
        for row in rows_from_sheet(workbook, "小料目录"):
            name = str(row["小料名称"]).strip()
            unit = clean_text(row.get("单位"))
            minimum = numeric(row.get("热量下限(kcal)"))
            maximum = numeric(row.get("热量上限(kcal)"))
            if minimum is None:
                minimum = numeric(row.get("参考热量(kcal)"))
            if maximum is None:
                maximum = minimum
            topping_candidates[name].append(
                make_topping_variant(
                    brand_id,
                    name,
                    None,
                    unit,
                    minimum,
                    maximum,
                    (row.get("参考份量"), row.get("适用说明"), row.get("内部备注")),
                )
            )
    toppings = group_toppings(brand_id, topping_candidates)
    topping_ids = [topping["toppingId"] for topping in toppings]

    products = []
    for row in catalog:
        product_name = str(row[name_field]).strip()
        product = product_record(
            brand_id,
            brand_name,
            row,
            name_field,
            category_field,
            variants_by_product.get(product_name, []),
            topping_ids,
            explicit_defaults.get(product_name),
            ("source-needs-review",) if clean_text(row.get("数据状态")) == "needs-review" else (),
        )
        supplied_category = clean_text(row.get("前台统一分类"))
        if supplied_category in DISPLAY_CATEGORIES:
            product["displayCategory"] = supplied_category
        product["notes"] = [
            text
            for value in (
                row.get("产品类型"),
                row.get("原始产品名称"),
                row.get("其他默认规格"),
                row.get("内部备注"),
            )
            if (text := clean_text(value)) is not None
        ]
        products.append(product)
    return {"brandId": brand_id, "brandName": brand_name, "products": products, "toppings": toppings}


def build_coco(workbook: Any, report: dict[str, Any]) -> dict[str, Any]:
    return build_normalized(workbook, report, "CoCo都可.xlsx", "热量矩阵")


def build_guming(workbook: Any, report: dict[str, Any]) -> dict[str, Any]:
    return build_normalized(workbook, report, "古茗.xlsx", "热量记录")


def build_luckin(workbook: Any, report: dict[str, Any]) -> dict[str, Any]:
    return build_normalized(workbook, report, "瑞幸咖啡.xlsx", "热量记录")


def build_chabaidao(workbook: Any, report: dict[str, Any]) -> dict[str, Any]:
    return build_normalized(workbook, report, "茶百道.xlsx", "热量矩阵")


def apply_editorial_overrides(brands: list[dict[str, Any]], report: dict[str, Any]) -> None:
    """Apply product-specific website decisions without changing source workbooks."""
    brand_by_id = {brand["brandId"]: brand for brand in brands}

    def product(brand_id: str, name: str) -> dict[str, Any]:
        matches = [item for item in brand_by_id[brand_id]["products"] if item["productName"] == name]
        if len(matches) != 1:
            raise ValueError(f"editorial override target must be unique: {brand_id} {name}")
        return matches[0]

    def selection_from_variant(variant: dict[str, Any]) -> dict[str, str | None]:
        return {field: variant[field] for field in SELECTION_FIELDS}

    def set_single_variant(
        item: dict[str, Any],
        calories: int | float,
        selection: dict[str, str | None] | None = None,
        status: str | None = None,
    ) -> None:
        exact_selection = selection or {field: None for field in SELECTION_FIELDS}
        variant = make_variant(item["brandId"], item["productName"], exact_selection, calories)
        if status is not None:
            variant["dataStatus"] = status
        variant["reviewFlags"] = []
        item["variants"] = [variant]
        item["defaultSelection"] = default_selection(variant)
        item["dataStatus"] = variant["dataStatus"]
        item["reviewFlags"] = []

    def set_default(item: dict[str, Any], variant: dict[str, Any]) -> None:
        item["defaultSelection"] = default_selection(variant)
        item["dataStatus"] = variant["dataStatus"]

    yidiandian = brand_by_id["yidiandian"]
    yidiandian["products"] = [item for item in yidiandian["products"] if item["productName"] != "奶香普洱"]
    if len(yidiandian["products"]) != 54:
        raise ValueError("一点点删除奶香普洱后的产品数应为 54")

    set_single_variant(product("yidiandian", "六窨香柠绿"), 5, status="unspecified_reference")
    set_single_variant(product("chagee", "龙井玄米酪"), 125, status="unspecified_reference")

    yidiandian_references = {
        "波霸奶茶": 400,
        "阿华田": 319,
        "A2牛乳红茶": 128,
        "奶绿装芒": 275,
        "QQ美莓奶茶": 300,
        "苹果奶绿": 247,
        "四季奶青": 184,
        "草莓A2牛乳绿茶": 205,
    }
    for name, calories in yidiandian_references.items():
        set_single_variant(product("yidiandian", name), calories, status="unspecified_reference")

    strawberry = product("chabaidao", "草莓流心半熟芝士")
    strawberry_calories = {"无糖": 263, "全糖": 278}
    for variant in strawberry["variants"]:
        if variant["sugar"] not in strawberry_calories:
            raise ValueError("草莓流心半熟芝士存在未预期甜度")
        variant["calories"] = strawberry_calories[variant["sugar"]]
        variant["dataStatus"] = data_status(selection_from_variant(variant), variant["calories"])
        variant["reviewFlags"] = []
    strawberry["reviewFlags"] = []
    set_default(strawberry, next(variant for variant in strawberry["variants"] if variant["sugar"] == "无糖"))

    mango = product("heytea", "椰椰芒芒")
    mango_calories = {"不另外加糖": 215, "少少少甜": 281, "少少甜": 364, "少甜": 377}
    mango_group = [variant for variant in mango["variants"] if variant["ice"] == "冰沙-少冰"]
    if {variant["sugar"] for variant in mango_group} != set(mango_calories):
        raise ValueError("椰椰芒芒冰沙-少冰甜度组不完整")
    for variant in mango_group:
        variant["calories"] = mango_calories[variant["sugar"]]
        variant["dataStatus"] = data_status(selection_from_variant(variant), variant["calories"])
        variant["reviewFlags"] = []

    matcha = product("chagee", "月抹静山")
    matcha_calories = {"不另外加糖": 360, "少糖": 432, "标准糖": 448}
    hot_group = [variant for variant in matcha["variants"] if variant["temperature"] == "热"]
    if {variant["sugar"] for variant in hot_group} != set(matcha_calories):
        raise ValueError("月抹静山热饮甜度组不完整")
    for variant in hot_group:
        variant["calories"] = matcha_calories[variant["sugar"]]
        variant["dataStatus"] = data_status(selection_from_variant(variant), variant["calories"])
        variant["reviewFlags"] = []
    matcha["reviewFlags"] = []

    fruit_tea = product("chabaidao", "超级杯水果茶")
    expected_fruit_tea = {"无糖": 275, "三分糖": 305, "五分糖": 319, "七分糖": 334}
    fruit_tea["variants"] = [
        variant
        for variant in fruit_tea["variants"]
        if variant["size"] == "大杯" and variant["ice"] == "标准冰" and variant["sugar"] in expected_fruit_tea
    ]
    if len(fruit_tea["variants"]) != 4:
        raise ValueError("超级杯水果茶应只保留四条大杯矩阵")
    for variant in fruit_tea["variants"]:
        variant["calories"] = expected_fruit_tea[variant["sugar"]]
        variant["dataStatus"] = data_status(selection_from_variant(variant), variant["calories"])
        variant["reviewFlags"] = []
    fruit_tea["reviewFlags"] = []
    set_default(fruit_tea, next(variant for variant in fruit_tea["variants"] if variant["sugar"] == "无糖"))

    yanhong = product("heytea", "嫣红牛乳茶")
    yanhong["reviewFlags"] = [flag for flag in yanhong["reviewFlags"] if flag != "catalog-mismatch"]
    yanhong_default = next(
        variant for variant in yanhong["variants"]
        if variant["calories"] == 255 and variant["ice"] == "冰" and variant["sugar"] == "不另外加糖"
    )
    set_default(yanhong, yanhong_default)

    luckin_matcha = product("luckin", "抹茶好喝椰")
    luckin_matcha["variants"] = [variant for variant in luckin_matcha["variants"] if variant["calories"] not in {149, 279}]
    luckin_matcha["reviewFlags"] = []
    for variant in luckin_matcha["variants"]:
        variant["reviewFlags"] = []
    matcha_default = next(
        variant for variant in luckin_matcha["variants"]
        if variant["calories"] == 273 and variant["ice"] == "冰" and variant["sugar"] == "不另外加糖"
    )
    set_default(luckin_matcha, matcha_default)

    heytea = brand_by_id["heytea"]
    regular_xiaonaimo = product("heytea", "小奶茉")
    oversized_xiaonaimo = product("heytea", "小奶茉（超大）")
    regular_default_id = regular_xiaonaimo["defaultSelection"]["variantId"]
    oversized_selection = selection_from_variant(oversized_xiaonaimo["variants"][0])
    oversized_selection["size"] = "超大杯"
    regular_xiaonaimo["variants"].append(
        make_variant("heytea", "小奶茉", oversized_selection, 202)
    )
    unique_xiaonaimo = {selection_key(variant): variant for variant in regular_xiaonaimo["variants"]}
    regular_xiaonaimo["variants"] = list(unique_xiaonaimo.values())
    regular_xiaonaimo["notes"] = [
        note for note in regular_xiaonaimo["notes"] if "小奶茉（超大）" not in note and "分开记录" not in note
    ]
    regular_default = next(variant for variant in regular_xiaonaimo["variants"] if variant["variantId"] == regular_default_id)
    set_default(regular_xiaonaimo, regular_default)
    heytea["products"] = [item for item in heytea["products"] if item["productName"] != "小奶茉（超大）"]

    butter_latte = product("luckin", "小黄油拿铁")
    butter_selection = {
        "size": "大杯",
        "ice": "正常冰",
        "temperature": None,
        "sugar": "不另外加糖",
        "version": None,
        "base": None,
    }
    set_single_variant(butter_latte, 282, butter_selection, status="specified_reference")

    luckin = brand_by_id["luckin"]
    luckin_food = [item for item in luckin["products"] if item["category"] == "简餐家族"]
    if len(luckin_food) != 49:
        raise ValueError(f"瑞幸简餐家族应为 49 款，实际 {len(luckin_food)}")
    for item in luckin["products"]:
        if item["category"] == "简餐家族":
            item["excludeFromWheel"] = True
        else:
            item.pop("excludeFromWheel", None)

    total_products = sum(len(brand["products"]) for brand in brands)
    if total_products != 515:
        raise ValueError(f"编辑覆盖后产品总数应为 515，实际 {total_products}")
    report["editorialOverrides"] = {
        "removedProducts": ["一点点 / 奶香普洱", "喜茶 / 小奶茉（超大）"],
        "luckinFoodExcludedFromWheel": len(luckin_food),
        "finalProductRules": [
            "一点点 / 六窨香柠绿：仅保留5 kcal",
            "霸王茶姬 / 龙井玄米酪：仅保留125 kcal",
            "茶百道 / 草莓流心半熟芝士：无糖263、全糖278 kcal",
            "喜茶 / 椰椰芒芒：仅修正冰沙-少冰糖度组",
            "霸王茶姬 / 月抹静山：仅修正热饮糖度组",
            "茶百道 / 超级杯水果茶：仅保留大杯标准冰四档糖度",
            "喜茶 / 嫣红牛乳茶：默认255 kcal",
            "瑞幸咖啡 / 抹茶好喝椰：采用273 kcal对应矩阵",
            "喜茶 / 小奶茉：并入超大杯202 kcal规格",
            "瑞幸咖啡 / 小黄油拿铁：仅保留大杯正常冰不另外加糖282 kcal",
        ],
    }


def scan_formula_errors(workbook: Any, filename: str) -> list[dict[str, str]]:
    errors = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                    errors.append({"file": filename, "sheet": worksheet.title, "cell": cell.coordinate, "value": value})
    return errors


def validate_brand(brand: dict[str, Any]) -> None:
    product_ids = [product["productId"] for product in brand["products"]]
    if len(product_ids) != len(set(product_ids)):
        raise ValueError(f"{brand['brandName']}: duplicate productId")
    product_names = [product["productName"] for product in brand["products"]]
    if len(product_names) != len(set(product_names)):
        raise ValueError(f"{brand['brandName']}: duplicate productName")
    for product in brand["products"]:
        if product["displayCategory"] not in DISPLAY_CATEGORIES:
            raise ValueError(f"{brand['brandName']} {product['productName']}: invalid display category")
        variant_ids = [variant["variantId"] for variant in product["variants"]]
        if not variant_ids or len(variant_ids) != len(set(variant_ids)):
            raise ValueError(f"{brand['brandName']} {product['productName']}: invalid variants")
        if product["defaultSelection"]["variantId"] not in variant_ids:
            raise ValueError(f"{brand['brandName']} {product['productName']}: invalid default")
        selection_keys = [selection_key(variant) for variant in product["variants"]]
        if len(selection_keys) != len(set(selection_keys)):
            raise ValueError(f"{brand['brandName']} {product['productName']}: duplicate selection")
        for variant in product["variants"]:
            calories = variant["calories"]
            if calories is not None and not isinstance(calories, (int, float)):
                raise ValueError(f"{brand['brandName']} {product['productName']}: invalid calories")


def main() -> None:
    project_root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=project_root.parent / "资料源" / "奶茶" / "各个奶茶品牌")
    parser.add_argument("--output-dir", type=Path, default=project_root / "data" / "milk-tea")
    arguments = parser.parse_args()

    builders = {
        "一点点.xlsx": build_yidiandian,
        "喜茶.xlsx": build_heytea,
        "蜜雪冰城.xlsx": build_mixue,
        "霸王茶姬.xlsx": build_chagee,
        "CoCo都可.xlsx": build_coco,
        "古茗.xlsx": build_guming,
        "瑞幸咖啡.xlsx": build_luckin,
        "茶百道.xlsx": build_chabaidao,
    }
    output_names = {brand_id: f"{brand_id}.json" for brand_id, _, _, _ in BRANDS.values()}
    report: dict[str, Any] = {"sourceFiles": [], "brands": {}, "skippedRecords": [], "formulaErrors": [], "sourceDecisions": {}}
    brands: list[dict[str, Any]] = []

    for filename, builder in builders.items():
        path = arguments.source_dir / filename
        if not path.exists():
            raise FileNotFoundError(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        report["sourceFiles"].append({"file": filename, "sheets": workbook.sheetnames})
        report["formulaErrors"].extend(scan_formula_errors(workbook, filename))
        if filename == "蜜雪冰城.xlsx":
            report["sourceDecisions"]["蜜雪冰城"] = rows_from_sheet(workbook, "筛选记录")
        elif filename == "霸王茶姬.xlsx":
            report["sourceDecisions"]["霸王茶姬"] = rows_from_sheet(workbook, "筛选与核验记录")
        brand = builder(workbook, report)
        validate_brand(brand)
        brands.append(brand)

    apply_editorial_overrides(brands, report)
    for brand in brands:
        validate_brand(brand)

    arguments.output_dir.mkdir(parents=True, exist_ok=True)
    brands_dir = arguments.output_dir / "brands"
    brands_dir.mkdir(parents=True, exist_ok=True)
    for brand in brands:
        path = brands_dir / output_names[brand["brandId"]]
        path.write_text(json.dumps(brand, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        variants = sum(len(product["variants"]) for product in brand["products"])
        review_variants = sum(bool(variant["reviewFlags"]) for product in brand["products"] for variant in product["variants"])
        report["brands"][brand["brandId"]] = {
            "brandName": brand["brandName"],
            "products": len(brand["products"]),
            "variants": variants,
            "toppings": len(brand["toppings"]),
            "reviewVariants": review_variants,
        }

    report["totals"] = {
        "products": sum(value["products"] for value in report["brands"].values()),
        "variants": sum(value["variants"] for value in report["brands"].values()),
        "toppings": sum(value["toppings"] for value in report["brands"].values()),
        "skippedRecords": len(report["skippedRecords"]),
        "formulaErrors": len(report["formulaErrors"]),
    }
    (arguments.output_dir / "import-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"valid": True, "output": str(arguments.output_dir), **report["totals"], "brands": report["brands"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
